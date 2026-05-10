from __future__ import annotations

import csv
from io import BytesIO

from app.portfolio import RebalanceRecommendation
from app.reports.export import generate_rebalance_csv, generate_rebalance_excel


def test_generate_rebalance_excel_creates_workbook() -> None:
    """测试生成的 Excel 是有效的。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000001",
            stock_name="平安银行",
            action="buy",
            target_weight=0.08,
            total_score=0.75,
            rank=1,
            reason="新增目标持仓",
            risk_note=None,
        ),
    ]

    output = generate_rebalance_excel(recommendations)

    assert isinstance(output, BytesIO)
    assert len(output.getvalue()) > 0


def test_generate_rebalance_excel_contains_headers() -> None:
    """测试 Excel 包含表头。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000001",
            stock_name="平安银行",
            action="buy",
            target_weight=0.08,
            total_score=0.75,
            rank=1,
            reason="新增目标持仓",
            risk_note=None,
        ),
    ]

    output = generate_rebalance_excel(recommendations)

    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(output.getvalue()))
    worksheet = workbook.active

    headers = ["股票代码", "股票名称", "操作", "目标仓位", "总分", "排名", "原因", "风险提示"]
    for col, header in enumerate(headers, start=1):
        assert worksheet.cell(1, col).value == header


def test_generate_rebalance_excel_buy_list() -> None:
    """测试买入列表正确渲染。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000001",
            stock_name="平安银行",
            action="buy",
            target_weight=0.08,
            total_score=0.75,
            rank=1,
            reason="新增目标持仓",
            risk_note=None,
        ),
    ]

    output = generate_rebalance_excel(recommendations)

    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(output.getvalue()))
    worksheet = workbook.active

    assert worksheet.cell(2, 1).value == "000001"
    assert worksheet.cell(2, 2).value == "平安银行"
    assert worksheet.cell(2, 3).value == "buy"
    assert worksheet.cell(2, 4).value == "8.00%"


def test_generate_rebalance_excel_sell_list() -> None:
    """测试卖出列表正确渲染。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000002",
            stock_name="万科A",
            action="sell",
            target_weight=None,
            total_score=0.60,
            rank=None,
            reason="已不在目标组合中",
            risk_note=None,
        ),
    ]

    output = generate_rebalance_excel(recommendations)

    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(output.getvalue()))
    worksheet = workbook.active

    assert worksheet.cell(2, 1).value == "000002"
    assert worksheet.cell(2, 2).value == "万科A"
    assert worksheet.cell(2, 3).value == "sell"


def test_generate_rebalance_csv_creates_csv() -> None:
    """测试生成的 CSV 是有效的。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000001",
            stock_name="平安银行",
            action="buy",
            target_weight=0.08,
            total_score=0.75,
            rank=1,
            reason="新增目标持仓",
            risk_note=None,
        ),
    ]

    csv_content = generate_rebalance_csv(recommendations)

    assert isinstance(csv_content, str)
    assert len(csv_content) > 0


def test_generate_rebalance_csv_contains_headers() -> None:
    """测试 CSV 包含表头。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000001",
            stock_name="平安银行",
            action="buy",
            target_weight=0.08,
            total_score=0.75,
            rank=1,
            reason="新增目标持仓",
            risk_note=None,
        ),
    ]

    csv_content = generate_rebalance_csv(recommendations)

    reader = csv.reader(csv_content.splitlines())
    headers = next(reader)

    expected_headers = [
        "股票代码",
        "股票名称",
        "操作",
        "目标仓位",
        "总分",
        "排名",
        "原因",
        "风险提示",
    ]
    assert headers == expected_headers


def test_generate_rebalance_csv_buy_list() -> None:
    """测试买入列表正确渲染。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000001",
            stock_name="平安银行",
            action="buy",
            target_weight=0.08,
            total_score=0.75,
            rank=1,
            reason="新增目标持仓",
            risk_note=None,
        ),
    ]

    csv_content = generate_rebalance_csv(recommendations)

    reader = csv.reader(csv_content.splitlines())
    next(reader)
    row = next(reader)

    assert row[0] == "000001"
    assert row[1] == "平安银行"
    assert row[2] == "buy"
    assert row[3] == "8.00%"


def test_generate_rebalance_csv_sell_list() -> None:
    """测试卖出列表正确渲染。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000002",
            stock_name="万科A",
            action="sell",
            target_weight=None,
            total_score=0.60,
            rank=None,
            reason="已不在目标组合中",
            risk_note=None,
        ),
    ]

    csv_content = generate_rebalance_csv(recommendations)

    reader = csv.reader(csv_content.splitlines())
    next(reader)
    row = next(reader)

    assert row[0] == "000002"
    assert row[1] == "万科A"
    assert row[2] == "sell"
    assert row[3] == "—"


def test_generate_rebalance_csv_risk_note() -> None:
    """测试风险提示正确渲染。"""
    recommendations = [
        RebalanceRecommendation(
            stock_code="000001",
            stock_name="平安银行",
            action="buy",
            target_weight=0.08,
            total_score=0.75,
            rank=1,
            reason="新增目标持仓",
            risk_note="涨停无法买入",
        ),
    ]

    csv_content = generate_rebalance_csv(recommendations)

    assert "涨停无法买入" in csv_content
